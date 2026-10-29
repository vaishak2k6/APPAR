# APAAR-Satya SIH 2026 idea deck: official SIH template, laid out and worded like the FDR deck.
# Run:  python build_human.py   ->  ../APAAR-Satya_SIH2026_Final.pptx
import json, os, subprocess
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree

HERE = os.path.dirname(os.path.abspath(__file__))
TEMPLATE = os.path.join(HERE, "tpl_a.pptx")
LOGO = os.path.join(HERE, "assets", "team_logo.png")
SHOTS = os.path.join(HERE, "shots")
OUT = os.path.join(os.path.dirname(HERE), "APAAR-Satya_SIH2026_Final.pptx")
ICON_DIR = os.path.join(HERE, "icons")

SERIF = "Times New Roman"
MONO = "Consolas"

# FDR-style palette: pastel gradient boxes, brown-red labels, black body text
INK = "111827"; LABEL = "953735"; NAVY = "0F172A"; INDIGO = "1E40AF"; EMERALD = "047857"; CRIMSON = "B91C1C"; AMBER = "B45309"
MINT = "6EE7B7"; MUTED = "4B5563"
PEACH = ("FDEBDD", "F8C9A4", "F4A261")      # top, bottom, border
LILAC = ("ECE4FB", "C3B1EC", "8E6FC9")
CYAN = ("E6F8FC", "A6E3F2", "4FB3D1")
GREEN = ("E6F4D7", "B9DD95", "6A9F3B")
SALMON = ("FDE3DA", "F6B6A0", "E07A5F")
WHITE = "FFFFFF"

ICONS = {
    "upload": "FaUpload", "idcard": "FaIdCard", "policy": "FaScaleBalanced", "shield": "FaShieldHalved",
    "db": "FaDatabase", "qr": "FaQrcode", "lock": "FaLock", "graph": "FaDiagramProject", "chart": "FaChartLine",
    "users": "FaUsers", "check": "FaCheck", "terminal": "FaTerminal", "route": "FaRoute", "plug": "FaPlug",
    "fingerprint": "FaFingerprint", "monitor": "FaDesktop", "flask": "FaFlask", "book": "FaBookOpen",
}


def render_icons():
    lst = []
    for k, v in ICONS.items():
        lst.append({"key": f"{k}_k", "name": v, "color": "1F2937"})
        lst.append({"key": f"{k}_w", "name": v, "color": "FFFFFF"})
    with open(os.path.join(HERE, "icons_fdr.json"), "w") as f:
        json.dump(lst, f)
    subprocess.run(["node", os.path.join(HERE, "icons.js"), os.path.join(HERE, "icons_fdr.json"), ICON_DIR], check=True)


def icon(key, v="k"):
    return os.path.join(ICON_DIR, f"{key}_{v}.png")


def rgb(h):
    return RGBColor.from_string(h)


# ---------- shape helpers ----------
def _clear_fill(spPr):
    for tag in ("a:solidFill", "a:noFill", "a:gradFill", "a:pattFill", "a:blipFill"):
        for el in spPr.findall(qn(tag)):
            spPr.remove(el)


def grad(shape, top, bottom, angle=90):
    spPr = shape._element.spPr
    _clear_fill(spPr)
    g = etree.SubElement(spPr, qn("a:gradFill")); g.set("rotWithShape", "1")
    gs = etree.SubElement(g, qn("a:gsLst"))
    for pos, c in ((0, top), (100000, bottom)):
        s = etree.SubElement(gs, qn("a:gs")); s.set("pos", str(pos))
        etree.SubElement(s, qn("a:srgbClr")).set("val", c)
    lin = etree.SubElement(g, qn("a:lin")); lin.set("ang", str(angle * 60000)); lin.set("scaled", "0")
    geom = spPr.find(qn("a:prstGeom"))
    geom.addnext(g)


def shadow(shape):
    spPr = shape._element.spPr
    for el in spPr.findall(qn("a:effectLst")):
        spPr.remove(el)
    eff = etree.SubElement(spPr, qn("a:effectLst"))
    sh = etree.SubElement(eff, qn("a:outerShdw"))
    sh.set("blurRad", "50800"); sh.set("dist", "25400"); sh.set("dir", "5400000"); sh.set("algn", "t"); sh.set("rotWithShape", "0")
    clr = etree.SubElement(sh, qn("a:srgbClr")); clr.set("val", "000000")
    etree.SubElement(clr, qn("a:alpha")).set("val", "22000")


def no_shadow(shape):
    spPr = shape._element.spPr
    if spPr.find(qn("a:effectLst")) is None:
        etree.SubElement(spPr, qn("a:effectLst"))


def box(s, x, y, w, h, fill=None, line=None, lw=1.25, kind=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.12, gradient=None, shade=False, dash=False):
    sh = s.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        sh.adjustments[0] = radius
    if fill:
        sh.fill.solid(); sh.fill.fore_color.rgb = rgb(fill)
    else:
        sh.fill.background()
    if gradient:
        grad(sh, gradient[0], gradient[1])
    if line:
        sh.line.color.rgb = rgb(line); sh.line.width = Pt(lw)
        if dash:
            sh.line.dash_style = 7
    else:
        sh.line.fill.background()
    if shade:
        shadow(sh)
    else:
        no_shadow(sh)
    sh.text_frame.text = ""
    return sh


def text(s, x, y, w, h, paras, size=13, color=INK, bold=False, italic=False, font=SERIF, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, margin=0.04, after=0, line_sp=None):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.auto_size = None
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{side}", Inches(margin))
    tf.vertical_anchor = anchor
    if isinstance(paras, (str, tuple)):
        paras = [paras]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if after:
            p.space_after = Pt(after)
        if line_sp:
            p.line_spacing = line_sp
        runs = [para] if isinstance(para, (str, tuple)) else para
        for r in runs:
            t, o = (r, {}) if isinstance(r, str) else r
            run = p.add_run(); run.text = t
            f = run.font
            f.size = Pt(o.get("size", size)); f.bold = o.get("bold", bold); f.italic = o.get("italic", italic)
            f.name = o.get("font", font); f.color.rgb = rgb(o.get("color", color))
    return tb


def pic(s, path, x, y, w, h=None):
    if h is None:
        return s.shapes.add_picture(path, Inches(x), Inches(y), Inches(w))
    return s.shapes.add_picture(path, Inches(x), Inches(y), Inches(w), Inches(h))


def arrow(s, x1, y1, x2, y2, color="1F2937", w=1.5, head=True):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = rgb(color); c.line.width = Pt(w)
    if head:
        ln = c.line._get_or_add_ln()
        t = etree.SubElement(ln, qn("a:tailEnd")); t.set("type", "triangle"); t.set("w", "med"); t.set("len", "med")
    return c


def pill(s, x, y, w, h, label, border, size=15, fill=WHITE, color=INK, kind=MSO_SHAPE.OVAL):
    sh = box(s, x, y, w, h, fill=fill, line=border, lw=1.75, kind=kind, radius=0.5)
    text(s, x, y, w, h, label, size=size, bold=True, color=color, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, margin=0.02)
    return sh


def items(s, x, y, w, h, rows, mark="❖", size=13, mark_color=LABEL, label_color=LABEL, after=3, font=SERIF):
    """FDR-style list: a mark, a coloured label, then plain text."""
    paras = []
    for lab, rest in rows:
        paras.append([(mark + " ", {"color": mark_color, "font": "Segoe UI Symbol" if mark != "✓" else SERIF}),
                      (lab, {"color": label_color, "bold": True}), (rest, {})])
    return text(s, x, y, w, h, paras, size=size, after=after, font=font)


def set_title(shape, txt, size=None):
    tf = shape.text_frame; p0 = tf.paragraphs[0]
    for extra in tf.paragraphs[1:]:
        extra._p.getparent().remove(extra._p)
    for br in p0._p.findall(qn("a:br")):
        p0._p.remove(br)
    runs = p0._p.findall(qn("a:r"))
    for r in runs[1:]:
        p0._p.remove(r)
    if runs:
        p0.runs[0].text = txt
    else:
        p0.add_run().text = txt
    if size:
        p0.runs[0].font.size = Pt(size)


def frame(slide, title, size=None):
    for sh in list(slide.shapes):
        if sh.name.startswith("TextBox 8"):
            sh._element.getparent().remove(sh._element)
        elif sh.name.startswith("Oval"):
            x, y, w, h = sh.left, sh.top, sh.width, sh.height
            sh._element.getparent().remove(sh._element)
            p = slide.shapes.add_picture(LOGO, x, y, w, h)
            c = max(0.0, (1 - (w / h) / (1024 / 559)) / 2)
            p.crop_left = c; p.crop_right = c
            p._element.spPr.find(qn("a:prstGeom")).set("prst", "ellipse")
            p.line.color.rgb = rgb("7030A0"); p.line.width = Pt(1.5)
        elif sh.is_placeholder and sh.placeholder_format.type == 1:
            set_title(sh, title, size)
    return slide


def delete_slide(prs, index):
    lst = prs.slides._sldIdLst
    sld = list(lst)[index]
    prs.part.drop_rel(sld.get(qn("r:id")))
    lst.remove(sld)


# Office bullet glyphs, the way PowerPoint itself inserts them: Wingdings v = ❖, ü = ✓, Ø = ➢
DIAMOND, TICK, ARROWHEAD = "v", "ü", "Ø"


def bullets(s, x, y, w, h, rows, glyph=DIAMOND, size=16, glyph_color=LABEL, label_color=LABEL, label_bold=True, after=6, caps=False):
    paras = []
    for lab, rest in rows:
        paras.append([(glyph, {"font": "Wingdings", "color": glyph_color}), ("  ", {}),
                      (lab.upper() if caps else lab, {"color": label_color, "bold": label_bold}), (rest, {})])
    return text(s, x, y, w, h, paras, size=size, after=after)


def framed_pic(s, path, x, y, w, h):
    p = pic(s, path, x, y, w, h)
    p.line.color.rgb = rgb("8C8C8C"); p.line.width = Pt(0.75)
    shadow(p)
    return p


# ======================================================================
render_icons()
prs = Presentation(TEMPLATE)
delete_slide(prs, 6)
S = prs.slides

# ---------------- 1. Title page: template fields only, like FDR ----------------
s1 = S[0]
tb = [sh for sh in s1.shapes if sh.name == "TextBox 9"][0]
body = tb._element.txBody
first = body.findall(qn("a:p"))[0]
if not first.findall(qn("a:r")):
    body.remove(first)
FIELDS = [
    ("Problem Statement ID", "Problem Statement ID –  ", "SIH26___"),
    ("Problem Statement Title", "Problem Statement Title-  ", "APAAR ID-Based Student Data Management and Verification Platform"),
    ("Theme", "Theme- ", "Smart Education"),
    ("PS Category", "PS Category- ", "Software"),
    ("Team ID", "Team ID- ", "______"),
    ("Team Name", "Team Name- ", "TEAM-RESILIENCE"),
]
for p in tb.text_frame.paragraphs:
    label = p.text.strip()
    for key, base, val in FIELDS:
        if label.startswith(key):
            for r in p.runs[1:]:
                r._r.getparent().remove(r._r)
            p.runs[0].text = base
            p.runs[0].font.size = Pt(22)
            r = p.add_run(); r.text = val
            r.font.bold = True; r.font.name = "Arial"; r.font.size = Pt(22)
            r.font.color.rgb = rgb("1F4E79")
            break
    pPr = p._p.get_or_add_pPr(); pPr.set("algn", "l")
    for old in pPr.findall(qn("a:lnSpc")) + pPr.findall(qn("a:spcBef")):
        pPr.remove(old)
    ln = etree.Element(qn("a:lnSpc")); etree.SubElement(ln, qn("a:spcPct")).set("val", "150000")
    pPr.insert(0, ln)
tb.top = Inches(2.05); tb.width = Inches(7.0)

# ---------------- 2. Idea ----------------
s2 = frame(S[1], "APAAR-SATYA", size=36)
t2 = [sh for sh in s2.shapes if sh.is_placeholder and sh.placeholder_format.type == 1][0]
t2.left, t2.top, t2.width, t2.height = Inches(2.0), Inches(0.0), Inches(8.5), Inches(0.75)
text(s2, 2.0, 0.66, 8.5, 0.4, "Student Data Management and Verification using APAAR ID", size=16, italic=True,
     color="1F4E79", align=PP_ALIGN.CENTER)

box(s2, 0.3, 1.38, 10.55, 1.95, gradient=PEACH[:2], line=PEACH[2], shade=True, radius=0.12)
pill(s2, 0.6, 1.17, 2.0, 0.4, "Core Issue", PEACH[2], size=16)
bullets(s2, 0.45, 1.62, 10.3, 1.7, [
    ("Only Totals Reach AICTE: ", "Most institutions report total enrolment, not student-level data."),
    ("Excess Admissions: ", "Seats filled beyond approved intake are hard to detect in time."),
    ("Quota Misuse and Ghost Students: ", "NRI / OCI / FN seats and fake entries are difficult to verify."),
    ("Duplicate Enrolment: ", "The same student may hold seats in two colleges."),
], size=16, after=4)
framed_pic(s2, os.path.join(SHOTS, "student_crop.jpg"), 11.2, 1.2, 1.55, 2.05)

box(s2, 0.3, 3.62, 7.05, 3.22, gradient=LILAC[:2], line=LILAC[2], shade=True, radius=0.12)
pill(s2, 0.6, 3.42, 2.9, 0.4, "Proposed Solution", LILAC[2], size=16)
bullets(s2, 0.45, 3.9, 6.85, 2.9, [
    ("APAAR Records: ", "Each admitted student is uploaded with APAAR ID."),
    ("Seat Ledger: ", "Seats issued as per approval; extra uploads are blocked."),
    ("Three-Source Check: ", "A record is verified only when three sources agree."),
    ("Fraud Check: ", "Duplicate APAAR IDs, similar names, shared mobiles."),
    ("Student Confirmation: ", "Students confirm via DigiLocker, SMS or IVR."),
    ("Verifiable Certificates: ", "QR code with APAAR ID and Student ID."),
], size=15, after=7)

box(s2, 7.55, 3.62, 5.5, 3.22, gradient=CYAN[:2], line=CYAN[2], shade=True, radius=0.12)
pill(s2, 9.35, 3.42, 1.9, 0.4, "Novelty", CYAN[2], size=16)
bullets(s2, 7.7, 3.9, 5.25, 2.9, [
    ("Truth Score: ", "Counselling, DigiLocker, NSP and ABC data are cross-checked."),
    ("Passive Liveness: ", "No semester credits means a possible ghost student."),
    ("Policy Simulator: ", "AICTE can test a new rule first."),
    ("Submit Once: ", "The same data fills AICTE, AISHE and NIRF reports."),
    ("Fair Review: ", "Dual-degree students are not penalised; colleges can appeal."),
], size=15, after=5)

# ---------------- 3. Technical approach ----------------
s3 = frame(S[2], "TECHNICAL APPROACH")
top = [
    ("Student Data Upload", "(Excel or API, signed by Principal)", CYAN, "upload"),
    ("APAAR Validation", "(ID format and duplicate check)", PEACH, "idcard"),
    ("Seat and Rule Check", "(approved intake, quota rules)", LILAC, "policy"),
    ("Multi-Source Verification", "(counselling, DigiLocker, NSP, ABC)", PEACH, "shield"),
    ("Central Repository", "(secure storage with audit trail)", CYAN, "db"),
]
bw, gap = 2.35, 0.25
centers = []
for i, (t, sub, pal, ic) in enumerate(top):
    x = 0.3 + i * (bw + gap)
    centers.append(x + bw / 2)
    box(s3, x, 1.2, bw, 1.0, gradient=pal[:2], line=pal[2], shade=True, radius=0.16)
    pic(s3, icon(ic), x + bw - 0.4, 1.3, 0.3, 0.3)
    text(s3, x + 0.1, 1.25, bw - 0.5, 0.95, [[(t, {"bold": False, "size": 15})], [(sub, {"size": 11.5})]], size=15)
    if i < 4:
        arrow(s3, x + bw + 0.02, 1.7, x + bw + gap - 0.02, 1.7, w=1.75)

cols = [
    ("Seat Ledger", CYAN, ["Seats issued from approval", "Extra uploads blocked", "Quota-wise tracking"]),
    ("Duplicate Check", PEACH, ["Hashed APAAR matching", "Similar name matching", "Shared mobile / bank links"]),
    ("Student Confirmation", LILAC, ["DigiLocker, SMS or IVR", "“I did not join” option", "Seat surrender request"]),
    ("Truth Score", PEACH, ["Score from 6 sources", "Verified or under review", "Reason for every flag"]),
    ("Verification", CYAN, ["QR on certificates", "Shows only what is needed", "Student sees access log"]),
]
arrow(s3, centers[2], 2.2, centers[2], 2.33, head=False)
arrow(s3, centers[0], 2.33, centers[-1], 2.33, head=False)
for i, (lab, pal, rows) in enumerate(cols):
    x = 0.3 + i * (bw + gap)
    arrow(s3, centers[i], 2.33, centers[i], 2.45)
    box(s3, x + 0.08, 2.45, bw - 0.16, 0.52, gradient=GREEN[:2], line=GREEN[2], kind=MSO_SHAPE.OVAL, shade=True)
    text(s3, x + 0.08, 2.45, bw - 0.16, 0.52, lab, size=14.5, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, margin=0.02)
    for j, r in enumerate(rows):
        y = 3.08 + j * 0.62
        box(s3, x, y, bw, 0.55, gradient=pal[:2], line=pal[2], radius=0.08)
        text(s3, x + 0.05, y, bw - 0.1, 0.55, r, size=13, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    arrow(s3, centers[i], 4.9, centers[i], 5.02, head=False)
arrow(s3, centers[0], 5.02, centers[-1], 5.02, head=False)
arrow(s3, 3.4, 5.02, 3.4, 5.13)
arrow(s3, 9.95, 5.02, 9.95, 5.13)

box(s3, 0.3, 5.13, 6.2, 0.78, gradient=SALMON[:2], line=SALMON[2], shade=True, radius=0.16)
pic(s3, icon("monitor"), 0.45, 5.33, 0.4, 0.4)
text(s3, 1.0, 5.13, 5.4, 0.78, [[("AICTE Dashboard", {"size": 15})], [("Alerts, fill rates, risky institutions and reports.", {"size": 12.5})]],
     anchor=MSO_ANCHOR.MIDDLE)
box(s3, 6.85, 5.13, 6.2, 0.78, gradient=("FDE7C8", "F7B36B"), line="E8913A", shade=True, radius=0.16)
pic(s3, icon("users"), 7.0, 5.33, 0.4, 0.4)
text(s3, 7.55, 5.13, 5.4, 0.78, [[("Officer Review and Decision", {"size": 15})], [("Authorised officer checks the evidence and decides.", {"size": 12.5})]],
     anchor=MSO_ANCHOR.MIDDLE)

box(s3, 0.3, 6.05, 12.75, 0.8, fill=WHITE, line="A6A6A6", lw=1, radius=0.12)
text(s3, 0.45, 6.05, 12.5, 0.8, [[("Technologies Used:  ", {"bold": True, "color": LABEL}),
                                  ("Next.js, React (frontend)  |  FastAPI, PostgreSQL (backend)  |  Python, NetworkX (fraud checks)  |  "
                                   "Keycloak, HashiCorp Vault (security)  |  Simulated APAAR, DigiLocker and NSP APIs", {})]],
     size=13.5, anchor=MSO_ANCHOR.MIDDLE)

# ---------------- 4. Feasibility and viability ----------------
s4 = frame(S[3], "FEASIBILITY AND VIABILITY")


def fv_card(s, x, title, pal, sections):
    w = 6.25
    box(s, x, 1.2, w, 5.65, gradient=pal[:2], line=pal[2], shade=True, radius=0.16)
    text(s, x, 1.28, w, 0.55, title, size=22, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    y = 1.92
    for k, (lab, body_txt) in enumerate(sections):
        h = 1.55
        box(s, x + 0.2, y, w - 0.4, h, fill=WHITE if k != 1 else pal[1], kind=MSO_SHAPE.RECTANGLE)
        text(s, x + 0.3, y + 0.06, w - 0.6, h - 0.1, [[(lab, {"bold": True, "color": LABEL, "size": 17})], body_txt], size=15)
        y += h + 0.03


fv_card(s4, 0.3, "FEASIBILITY", CYAN, [
    ("Analysis:", "Buildable with standard web technologies such as Next.js, FastAPI and PostgreSQL. Prototype screens and core checks are already built and tested on synthetic student records."),
    ("Challenges:", "No direct access to APAAR, counselling and NSP data yet. Spelling differences in Indian names may create false duplicate matches."),
    ("Strategies:", "Use simulated APIs with the real data formats. Normalise common name variants and send doubtful cases for manual review."),
])
fv_card(s4, 6.8, "VIABILITY", PEACH, [
    ("Analysis:", "Low-cost platform for AICTE. Colleges save time because one verified upload fills several reports."),
    ("Challenges:", "Real data access needs government approval. College adoption and DPDP Act compliance are also challenges."),
    ("Strategies:", "Start with one state and its counselling cell. Use role-based access, consent records and audit logs. Free for colleges; paid verification for employers."),
])

# ---------------- 5. Impact and benefits ----------------
s5 = frame(S[4], "IMPACT AND BENEFITS")
box(s5, 0.4, 1.68, 6.15, 5.12, gradient=CYAN[:2], line=CYAN[2], shade=True, radius=0.16)
pill(s5, 1.1, 1.2, 4.0, 0.72, "POTENTIAL\nIMPACTS", CYAN[2], size=17, fill=CYAN[0], color=LABEL)
bullets(s5, 0.6, 2.15, 5.8, 4.6, [
    ("Accurate Records: ", "AICTE gets student-level data instead of only total numbers."),
    ("Excess Admission Control: ", "Seats beyond approved intake are blocked at upload."),
    ("Scholarship Protection: ", "Scholarships reach only verified students."),
    ("Faster Inspections: ", "Risk-based sampling instead of visiting every college."),
    ("Student Protection: ", "Students can report fake enrolment and surrender seats."),
    ("Better Planning: ", "Real fill-rates help AICTE decide next year’s approvals."),
], glyph=TICK, size=17.5, glyph_color=INK, label_color=INK, label_bold=False, after=8)

box(s5, 6.85, 1.68, 6.1, 5.12, gradient=PEACH[:2], line=PEACH[2], shade=True, radius=0.16)
pill(s5, 8.1, 1.2, 3.6, 0.6, "BENEFITS", PEACH[2], size=17, fill=PEACH[0], color=LABEL)
bullets(s5, 7.05, 2.15, 5.75, 4.6, [
    ("Less Paperwork: ", "One verified upload fills AICTE, AISHE and NIRF data."),
    ("Trusted Certificates: ", "Employers verify degrees by scanning a QR code."),
    ("Privacy: ", "Caste, quota and marks are not shown to verifiers."),
    ("Transparency: ", "Every change is saved in a tamper-evident audit log."),
    ("Sustainable Model: ", "Free for students and colleges; employers pay for verification."),
    ("Rural Reach: ", "Students without smartphones confirm by SMS or voice call."),
], glyph=TICK, size=17.5, glyph_color=INK, label_color=INK, label_bold=False, after=8)

# ---------------- 6. Research and references ----------------
s6 = frame(S[5], "RESEARCH AND REFERENCES")
refs_l = [
    "[1] Ministry of Education (PIB), APAAR backgrounder, June 2026: 26.30 crore APAAR IDs generated. Available: static.pib.gov.in",
    "[2] I. Fellegi and A. Sunter, “A Theory for Record Linkage,” JASA, vol. 64, 1969. doi.org/10.1080/01621459.1969.10501049",
    "[3] L. Akoglu, H. Tong and D. Koutra, “Graph-based anomaly detection and description: a survey,” DMKD, 2015.",
]
refs_r = [
    "[4] AICTE Approval Process Handbook 2026–27 and UGC guidelines for pursuing two academic programmes simultaneously, 2022.",
    "[5] B. Laurie et al., RFC 6962: Certificate Transparency (Merkle audit logs), IETF, 2013.",
    "[6] W3C Verifiable Credentials Data Model 2.0; Digital Personal Data Protection Act, 2023.",
]
for i, r in enumerate(refs_l):
    box(s6, 0.3, 1.2 + i * 0.52, 6.4, 0.5, fill=["8BE0F2", "B9ECF7", "8BE0F2"][i], kind=MSO_SHAPE.RECTANGLE)
    text(s6, 0.35, 1.2 + i * 0.52, 6.3, 0.5, r, size=11.5, anchor=MSO_ANCHOR.MIDDLE)
for i, r in enumerate(refs_r):
    box(s6, 6.85, 1.2 + i * 0.52, 6.2, 0.5, fill=["F8B67A", "FBD3AE", "F8B67A"][i], kind=MSO_SHAPE.RECTANGLE)
    text(s6, 6.9, 1.2 + i * 0.52, 6.1, 0.5, r, size=11.5, anchor=MSO_ANCHOR.MIDDLE)

box(s6, 0.3, 2.9, 4.75, 2.15, gradient=LILAC[:2], line=LILAC[2], shade=True, radius=0.14)
text(s6, 0.45, 2.94, 4.5, 0.32, "SOLUTIONS ALREADY EXIST", size=14, bold=True, color=LABEL)
bullets(s6, 0.45, 3.28, 4.5, 1.75, [
    ("", "AICTE portal collects only total intake."),
    ("", "AISHE survey collects self-reported statistics."),
    ("", "DigiLocker and NAD store issued certificates."),
    ("", "APAAR and ABC give each student an ID."),
], glyph=ARROWHEAD, size=14.5, glyph_color=INK, after=5)

box(s6, 5.2, 2.9, 7.85, 2.15, gradient=LILAC[:2], line=LILAC[2], shade=True, radius=0.14)
text(s6, 5.35, 2.94, 7.6, 0.32, "OUR SOLUTION STANDS OUT", size=14, bold=True, color=LABEL)
bullets(s6, 5.35, 3.28, 7.6, 1.75, [
    ("ONE VERIFIED RECORD: ", "links APAAR, counselling, DigiLocker, NSP, ABC."),
    ("SEAT LEDGER: ", "blocks admissions beyond approved intake at upload."),
    ("STUDENT CONFIRMATION: ", "students confirm or report their own admission."),
    ("EXPLAINABLE FLAGS: ", "every flag shows its reason for officer review."),
    ("PROTOTYPE READY: ", "dashboards and checks shown on synthetic data."),
], glyph=ARROWHEAD, size=14.5, glyph_color=INK, label_color=INK, after=3)

for i, f in enumerate(["overview.jpg", "upload.jpg", "institution.jpg", "verify.jpg"]):
    framed_pic(s6, os.path.join(SHOTS, f), 0.3 + i * 2.42, 5.25, 2.3, 1.45)
box(s6, 10.15, 5.18, 2.9, 1.6, gradient=GREEN[:2], line=GREEN[2], kind=MSO_SHAPE.OVAL, shade=True)
text(s6, 10.3, 5.28, 2.6, 1.4, [[("Link:", {})], [("[add deployed link]", {"color": "1F4E79", "underline": True})],
                               [("GitHub link:", {})], [("[add repo link]", {"color": "1F4E79"})]],
     size=13, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

prs.save(OUT)
print("saved", OUT)


